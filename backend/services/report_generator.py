"""
report_generator.py
Beautiful PDF/DOCX/CSV/XLSX reports from chat messages including charts.
Messages format: [{role, content, chart:{type,title,labels,values}, table:{columns,rows}, followup}]
"""

import io
from datetime import datetime
from typing import List, Dict, Any
import random, string


def _ref_id():
    return ''.join(random.choices(string.ascii_uppercase + string.digits, k=8))


# ─── PDF ─────────────────────────────────────────────────────
def generate_pdf(messages: List[Dict[str, Any]]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                     Table, TableStyle, HRFlowable)
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
        rightMargin=18*mm, leftMargin=18*mm, topMargin=16*mm, bottomMargin=16*mm)

    W = A4[0] - 36*mm
    ref = _ref_id()
    now = datetime.now().strftime("%-d/%-m/%Y, %-I:%M:%S %p") if hasattr(datetime, 'now') else datetime.now().strftime("%d/%m/%Y, %I:%M:%S %p")

    # Styles
    styles = getSampleStyleSheet()
    def ps(name, **kw):
        return ParagraphStyle(name, parent=styles['Normal'], **kw)

    hdr_style    = ps('hdr', fontSize=22, fontName='Helvetica-Bold', textColor=colors.white)
    meta_style   = ps('meta', fontSize=9,  textColor=colors.white)
    section_style= ps('sec', fontSize=13, fontName='Helvetica-Bold', textColor=colors.HexColor('#1a1a2e'), spaceBefore=14, spaceAfter=4)
    query_style  = ps('qry', fontSize=11, fontName='Helvetica-Oblique', textColor=colors.HexColor('#333366'), spaceAfter=6)
    body_style   = ps('bdy', fontSize=10, leading=16, textColor=colors.HexColor('#222222'), spaceAfter=4)
    bullet_style = ps('bul', fontSize=10, leading=16, textColor=colors.HexColor('#222222'), leftIndent=14, spaceAfter=2)
    followup_style=ps('flw', fontSize=9, fontName='Helvetica-Oblique', textColor=colors.HexColor('#555577'), spaceAfter=8)
    user_style   = ps('usr', fontSize=11, fontName='Helvetica-Bold', textColor=colors.HexColor('#4b0082'), spaceBefore=10, spaceAfter=2)

    story = []

    # ── Header banner ──
    header_data = [[
        Paragraph("InsightAI Analysis Report", hdr_style),
        Paragraph(f"REFERENCE ID: {ref}<br/>GENERATED: {datetime.now().strftime('%d/%m/%Y, %I:%M:%S %p')}", meta_style)
    ]]
    ht = Table(header_data, colWidths=[W*0.6, W*0.4])
    ht.setStyle(TableStyle([
        ('BACKGROUND', (0,0),(-1,-1), colors.HexColor('#1a1a2e')),
        ('VALIGN',     (0,0),(-1,-1), 'MIDDLE'),
        ('ALIGN',      (1,0),(1,0),   'RIGHT'),
        ('TOPPADDING', (0,0),(-1,-1), 14),
        ('BOTTOMPADDING',(0,0),(-1,-1),14),
        ('LEFTPADDING',(0,0),(0,0),   14),
        ('RIGHTPADDING',(1,0),(1,0),  14),
    ]))
    story.append(ht)
    story.append(Spacer(1, 10))

    # ── Messages ──
    qa_pairs = []
    i = 0
    msg_list = [m for m in messages if m.get('role') in ('user','assistant')]
    while i < len(msg_list):
        if msg_list[i].get('role') == 'user':
            q = msg_list[i]
            a = msg_list[i+1] if i+1 < len(msg_list) and msg_list[i+1].get('role') == 'assistant' else None
            qa_pairs.append((q, a))
            i += 2
        else:
            i += 1

    for idx, (q, a) in enumerate(qa_pairs, 1):
        # Query box
        story.append(Paragraph(f"ANALYTICAL QUERY", section_style))
        story.append(Paragraph(q.get('content',''), query_style))
        story.append(HRFlowable(width=W, thickness=0.5, color=colors.HexColor('#ccccdd')))
        story.append(Spacer(1, 6))

        if a:
            story.append(Paragraph("EXECUTIVE SYNTHESIS", section_style))
            content = a.get('content', '')
            # Split bullets
            lines = content.split('\n')
            for line in lines:
                line = line.strip()
                if not line: continue
                safe = line.replace('<','&lt;').replace('>','&gt;')
                if line.startswith('•'):
                    story.append(Paragraph(safe, bullet_style))
                else:
                    story.append(Paragraph(safe, body_style))

            # ── Chart ──
            chart = a.get('chart')
            if chart and chart.get('type') not in (None, 'none') and chart.get('labels') and chart.get('values'):
                ctype  = chart.get('type','bar')
                labels = chart.get('labels', [])
                values = [float(v) for v in chart.get('values', [])]
                title  = chart.get('title','Chart')

                story.append(Spacer(1, 8))
                story.append(Paragraph(title, ps('ct', fontSize=10, fontName='Helvetica-Bold',
                    textColor=colors.HexColor('#1a1a2e'), spaceAfter=4)))

                COLORS = [colors.HexColor(c) for c in [
                    '#4b0082','#7b1fa2','#e53935','#1565c0','#2e7d32',
                    '#e65100','#00838f','#f9a825','#ad1457','#37474f'
                ]]

                if ctype == 'pie' and len(labels) > 0:
                    d = Drawing(W, 180)
                    pie = Pie()
                    pie.x = W//2 - 70; pie.y = 10
                    pie.width = 140; pie.height = 140
                    pie.data  = values[:10]
                    pie.labels= [str(l)[:15] for l in labels[:10]]
                    pie.sideLabels = True
                    for j in range(len(pie.data)):
                        pie.slices[j].fillColor = COLORS[j % len(COLORS)]
                    d.add(pie)
                    story.append(d)

                elif ctype in ('bar', 'line') and len(values) > 0:
                    d = Drawing(W, 160)
                    bc = VerticalBarChart()
                    bc.x = 40; bc.y = 20
                    bc.width  = W - 60
                    bc.height = 120
                    bc.data   = [values[:12]]
                    bc.categoryAxis.categoryNames = [str(l)[:10] for l in labels[:12]]
                    bc.bars[0].fillColor = colors.HexColor('#4b0082')
                    bc.valueAxis.visibleGrid = True
                    bc.categoryAxis.labels.angle = 30 if len(labels) > 5 else 0
                    bc.categoryAxis.labels.fontSize = 7
                    d.add(bc)
                    story.append(d)

                story.append(Spacer(1, 8))

            # ── Table ──
            tbl = a.get('table')
            if tbl and tbl.get('columns') and tbl.get('rows'):
                cols = tbl['columns']
                rows = tbl['rows']
                story.append(Spacer(1, 6))
                col_w = W / max(len(cols), 1)
                tdata = [cols] + [[str(c) for c in r] for r in rows[:30]]
                t = Table(tdata, colWidths=[col_w]*len(cols))
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0),(-1,0), colors.HexColor('#4b0082')),
                    ('TEXTCOLOR',  (0,0),(-1,0), colors.white),
                    ('FONTNAME',   (0,0),(-1,0), 'Helvetica-Bold'),
                    ('FONTSIZE',   (0,0),(-1,-1), 8),
                    ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#f8f7ff'), colors.white]),
                    ('GRID',       (0,0),(-1,-1), 0.3, colors.HexColor('#ccccdd')),
                    ('TOPPADDING', (0,0),(-1,-1), 4),
                    ('BOTTOMPADDING',(0,0),(-1,-1),4),
                ]))
                story.append(t)
                story.append(Spacer(1, 6))

            # followup
            fu = a.get('followup','')
            if fu:
                story.append(Paragraph(f"💡 {fu}", followup_style))

        story.append(Spacer(1, 14))
        if idx < len(qa_pairs):
            story.append(HRFlowable(width=W, thickness=1, color=colors.HexColor('#e0e0f0')))
            story.append(Spacer(1, 8))

    doc.build(story)
    return buf.getvalue()


# ─── DOCX ────────────────────────────────────────────────────
def generate_docx(messages: List[Dict[str, Any]]) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    import struct, zlib

    doc = Document()
    # Page margins
    for sec in doc.sections:
        sec.top_margin    = Cm(1.8)
        sec.bottom_margin = Cm(1.8)
        sec.left_margin   = Cm(2)
        sec.right_margin  = Cm(2)

    def add_colored_heading(text, color_hex='1a1a2e', size=18):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(
            int(color_hex[:2],16), int(color_hex[2:4],16), int(color_hex[4:6],16))
        return p

    def add_body(text, italic=False, color_hex='222222', size=10):
        p = doc.add_paragraph()
        for line in text.split('\n'):
            line = line.strip()
            if not line: continue
            run = p.add_run(('• ' if line.startswith('•') else '') + line.lstrip('•').strip())
            run.font.size = Pt(size)
            run.font.italic = italic
            run.font.color.rgb = RGBColor(
                int(color_hex[:2],16), int(color_hex[2:4],16), int(color_hex[4:6],16))
        return p

    # Title
    add_colored_heading("InsightAI Analysis Report", '1a1a2e', 20)
    p = doc.add_paragraph(f"Generated: {datetime.now().strftime('%d/%m/%Y %I:%M %p')}  |  Ref: {_ref_id()}")
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor(0x77,0x77,0x99)
    doc.add_paragraph('─'*60)

    msg_list = [m for m in messages if m.get('role') in ('user','assistant')]
    i = 0
    while i < len(msg_list):
        if msg_list[i].get('role') == 'user':
            q = msg_list[i]
            a = msg_list[i+1] if i+1 < len(msg_list) else None
            i += 2

            add_colored_heading("ANALYTICAL QUERY", '4b0082', 12)
            add_body(q.get('content',''), italic=True, color_hex='333366', size=11)
            doc.add_paragraph()

            if a:
                add_colored_heading("EXECUTIVE SYNTHESIS", '1a1a2e', 12)
                add_body(a.get('content',''), size=10)

                # Table
                tbl = a.get('table')
                if tbl and tbl.get('columns') and tbl.get('rows'):
                    cols = tbl['columns']
                    rows = tbl['rows'][:30]
                    t = doc.add_table(rows=1+len(rows), cols=len(cols))
                    t.style = 'Table Grid'
                    hdr = t.rows[0].cells
                    for j, c in enumerate(cols):
                        hdr[j].text = str(c)
                        hdr[j].paragraphs[0].runs[0].bold = True
                    for r, row in enumerate(rows):
                        for j, cell in enumerate(row):
                            t.rows[r+1].cells[j].text = str(cell)

                fu = a.get('followup','')
                if fu:
                    p = doc.add_paragraph(f"💡 Suggested: {fu}")
                    p.runs[0].font.italic = True
                    p.runs[0].font.size = Pt(9)

            doc.add_paragraph('─'*60)
        else:
            i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─── CSV ─────────────────────────────────────────────────────
def generate_csv(messages: List[Dict[str, Any]]) -> bytes:
    import csv
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["#","Role","Message","Chart Type","Chart Labels","Chart Values","Followup"])
    for i, msg in enumerate(messages, 1):
        if msg.get('role') not in ('user','assistant'): continue
        chart  = msg.get('chart') or {}
        labels = '|'.join(str(l) for l in chart.get('labels',[])) if chart else ''
        values = '|'.join(str(v) for v in chart.get('values',[])) if chart else ''
        ctype  = chart.get('type','') if chart else ''
        w.writerow([i, msg['role'].upper(),
                    msg.get('content','').replace('\n',' '),
                    ctype, labels, values,
                    msg.get('followup','')])
    return buf.getvalue().encode('utf-8')


# ─── EXCEL ───────────────────────────────────────────────────
def generate_excel(messages: List[Dict[str, Any]]) -> bytes:
    import xlsxwriter
    buf = io.BytesIO()
    wb  = xlsxwriter.Workbook(buf, {'in_memory': True})

    # Formats
    hdr_fmt  = wb.add_format({'bold':True,'bg_color':'#1a1a2e','font_color':'white','border':1,'font_size':11})
    user_fmt = wb.add_format({'bg_color':'#f0efff','text_wrap':True,'valign':'top','border':1})
    ai_fmt   = wb.add_format({'bg_color':'#f0fff4','text_wrap':True,'valign':'top','border':1})
    num_fmt  = wb.add_format({'align':'center','valign':'top','border':1})
    title_fmt= wb.add_format({'bold':True,'font_size':16,'font_color':'#1a1a2e'})
    meta_fmt = wb.add_format({'font_size':9,'font_color':'#777799','italic':True})

    # Sheet 1: Chat
    ws = wb.add_worksheet("Chat Report")
    ws.set_column(0,0,5); ws.set_column(1,1,12); ws.set_column(2,2,70); ws.set_column(3,3,30)
    ws.write(0,0,"InsightAI Chat Report", title_fmt)
    ws.write(1,0,f"Generated: {datetime.now().strftime('%d/%m/%Y %I:%M %p')}", meta_fmt)
    ws.write_row(3,0,["#","Role","Message","Followup"], hdr_fmt)
    row = 4
    for i, msg in enumerate(messages,1):
        if msg.get('role') not in ('user','assistant'): continue
        fmt = user_fmt if msg['role']=='user' else ai_fmt
        ws.write(row,0,i,num_fmt)
        ws.write(row,1,msg['role'].upper(),fmt)
        ws.write(row,2,msg.get('content',''),fmt)
        ws.write(row,3,msg.get('followup',''),fmt)
        ws.set_row(row, max(30, min(120, len(msg.get('content',''))//2)))
        row += 1

    # Sheet 2: Charts data
    ws2 = wb.add_worksheet("Charts Data")
    ws2.set_column(0,0,20); ws2.set_column(1,1,50); ws2.set_column(2,2,50)
    ws2.write_row(0,0,["Chart Title","Labels","Values"], hdr_fmt)
    r2 = 1
    for msg in messages:
        chart = msg.get('chart')
        if chart and chart.get('type') not in (None,'none') and chart.get('labels'):
            ws2.write(r2,0,chart.get('title',''))
            ws2.write(r2,1,' | '.join(str(l) for l in chart.get('labels',[])))
            ws2.write(r2,2,' | '.join(str(v) for v in chart.get('values',[])))
            r2 += 1

    wb.close()
    return buf.getvalue()